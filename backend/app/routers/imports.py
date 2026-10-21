"""Импорт контактов из CSV/XLSX."""

from __future__ import annotations

import csv
import io
import logging
from typing import Annotated, Any

from fastapi import APIRouter, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select

from ..db import session_scope
from ..models import (
    Client,
    School,
    SchoolClass,
)
from ..phones import normalize_phone
from ..schemas import (
    ContactImport,
)
from ..services import (
    detect_header_row,
)

log = logging.getLogger("fotomaniabot")
router = APIRouter()

MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def _rows_from_upload(filename: str, raw: bytes) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        for merged_range in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            if min_col != max_col:
                continue
            value = sheet.cell(min_row, min_col).value
            if value is None:
                continue
            column_index = min_col - 1
            for row_index in range(min_row, max_row + 1):
                rows[row_index - 1][column_index] = value
        workbook.close()
        return rows
    text = raw.decode("utf-8-sig", errors="replace")
    counts = {candidate: text.count(candidate) for candidate in ",;\t|"}
    delimiter = (
        max(counts, key=lambda candidate: counts[candidate]) if any(counts.values()) else ","
    )
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _build_preview(rows: list[list[Any]]) -> dict[str, Any]:
    def cell_text(cell: Any) -> str:
        if cell is None:
            return ""
        # Excel returns whole numbers as float (89099695363.0); show them as integers.
        if isinstance(cell, float) and cell.is_integer():
            return str(int(cell))
        return str(cell).strip()

    header_index = detect_header_row(rows)
    raw_headers = rows[header_index] if rows else []
    headers = [cell_text(cell) for cell in raw_headers]
    data_rows = rows[header_index + 1 :]
    preview_rows = []
    for row in data_rows:
        if all(cell_text(cell) == "" for cell in row):
            continue
        record = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            record[header] = "" if index >= len(row) else cell_text(row[index])
        preview_rows.append(record)
    return {
        "header_row": header_index + 1,
        "headers": headers,
        "rows": preview_rows,
        "total_rows": len(preview_rows),
    }


@router.post("/api/imports/preview")
async def import_preview(file: Annotated[UploadFile, File()]) -> dict[str, Any]:
    raw = await file.read()
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Файл слишком большой (максимум 10 МБ)")
    try:
        rows = _rows_from_upload(file.filename or "upload.csv", raw)
    except Exception as exc:
        raise HTTPException(400, f"Не удалось прочитать файл: {exc}") from exc
    return _build_preview(rows)


def _resolve_school(session: Any, school_name: str) -> School:
    """Accept either a numeric school id or a human-readable school name."""
    value = (school_name or "").strip()
    if not value:
        raise HTTPException(400, "Не указано название школы")
    if value.isdigit():
        school = session.get(School, int(value))
        if school is not None:
            return school
    school = session.scalar(select(School).where(School.name == value))
    if school is None:
        school = School(name=value)
        session.add(school)
        session.flush()
    return school


def _resolve_class(session: Any, school_id: int, class_name: str) -> SchoolClass:
    value = (class_name or "").strip()
    if not value:
        value = "Без класса"
    item = session.scalar(
        select(SchoolClass).where(SchoolClass.school_id == school_id, SchoolClass.name == value)
    )
    if item is None:
        item = SchoolClass(school_id=school_id, name=value)
        session.add(item)
        session.flush()
    return item


def _normalize_import_header(value: Any) -> str:
    return "".join(str(value or "").casefold().replace("ё", "е").split())


def _import_row_value(row: dict[str, Any], header: str | None) -> Any:
    if header is None:
        return None
    if header in row:
        return row[header]
    normalized_header = _normalize_import_header(header)
    return next(
        (value for key, value in row.items() if _normalize_import_header(key) == normalized_header),
        None,
    )


@router.post("/api/imports/school-contacts")
def import_school_contacts(data: ContactImport) -> dict[str, int]:
    if not data.school_name:
        raise HTTPException(400, "Не указано название школы")
    created = updated = skipped = 0
    headers_by_field: dict[str, str] = {}
    for header, field in data.mapping.items():
        if not field:
            continue
        # При нескольких столбцах с одним полем берём первый по порядку в файле.
        headers_by_field.setdefault(field, header)
    class_header = headers_by_field.get("class_name")
    last_class_name = ""
    with session_scope() as session:
        school = _resolve_school(session, data.school_name)
        for row in data.rows:
            class_name = (
                str(_import_row_value(row, class_header) or "").strip() if class_header else ""
            )
            if class_name:
                last_class_name = class_name
            elif last_class_name:
                # Excel often stores the class only in the first row of a merged cell.
                class_name = last_class_name
            item = _resolve_class(session, school.id, class_name)
            for role, name_field, phone_field in (
                ("KR", "kr_name", "kr_phone"),
                ("RK", "rk_name", "rk_phone"),
            ):
                name_header = headers_by_field.get(name_field)
                if name_header is None:
                    continue
                name = str(_import_row_value(row, name_header) or "").strip()
                if not name:
                    continue
                phone_header = headers_by_field.get(phone_field)
                phone = normalize_phone(
                    str(_import_row_value(row, phone_header) or "").strip()
                    if phone_header
                    else None
                )
                duplicate = session.scalar(
                    select(Client).where(
                        Client.school_id == school.id,
                        Client.class_id == item.id,
                        Client.name == name,
                        Client.role == role,
                    )
                )
                if duplicate is None:
                    duplicate = session.scalar(
                        select(Client)
                        .outerjoin(SchoolClass, Client.class_id == SchoolClass.id)
                        .where(
                            Client.school_id == school.id,
                            Client.name == name,
                            Client.role == role,
                            (Client.class_id.is_(None)) | (SchoolClass.name == "Без класса"),
                        )
                    )
                if duplicate is not None:
                    if data.duplicate_mode == "update" and duplicate.class_id != item.id:
                        duplicate.class_id = item.id
                        duplicate.phone = phone or duplicate.phone
                        updated += 1
                    else:
                        skipped += 1
                    continue
                session.add(
                    Client(
                        name=name,
                        phone=phone,
                        role=role,
                        school_id=school.id,
                        class_id=item.id,
                        source="import",
                    )
                )
                created += 1
        session.commit()
    return {"created": created, "updated": updated, "skipped": skipped}
