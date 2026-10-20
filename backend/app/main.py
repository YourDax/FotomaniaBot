"""FastAPI-приложение: маршруты, воркер и запуск.

Модели, схемы, доступ к БД и прикладные функции вынесены в отдельные
модули (`models`, `schemas`, `db`, `services`, `security`), поэтому этот
файл содержит только маршруты и оркестрацию.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import threading
import time
import uuid
from contextlib import asynccontextmanager
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any, AsyncIterator

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from .db import BASE_DIR, engine, session_scope
from .max_browser import MaxBrowserError, open_chat, prepare_message
from .migrations import apply_migrations
from .models import (
    Base,
    Campaign,
    Client,
    DeliveryTask,
    School,
    SchoolClass,
    Setting,
    now,
)
from .phones import normalize_phone
from .schemas import (
    CampaignDeleteIn,
    CampaignIn,
    ClassIn,
    ClassUpdateIn,
    ClientDeleteIn,
    ClientIn,
    ContactImport,
    ImportConfirm,
    LoginIn,
    ManualResultIn,
    SchoolUpdateIn,
    SettingsIn,
    TaskStatusIn,
)
from .security import hash_password, require_auth
from .services import (
    client_dict,
    csv_safe,
    detect_header_row,
    render_message,
)

# Обратная совместимость: тесты обращаются к `main.db`
from .db import db  # noqa: E402  (re-export)

log = logging.getLogger("fotomaniabot")
LOG_DIR = BASE_DIR / "logs"
LOG_DIR.mkdir(exist_ok=True)
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s"
)
if not any(isinstance(handler, logging.FileHandler) for handler in log.handlers):
    file_handler = RotatingFileHandler(
        LOG_DIR / "session.log",
        maxBytes=5_000_000,
        backupCount=5,
        encoding="utf-8",
    )
    file_handler.setFormatter(
        logging.Formatter("%(asctime)s %(levelname)s %(name)s %(message)s")
    )
    log.addHandler(file_handler)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
DUPLICATE_FIELDS = ("max_user_id", "username", "chat_link")
MAX_SNAPSHOTS = 20
HEALTH_PATH = "/api/health"

worker_lock = threading.Lock()
worker_thread: threading.Thread | None = None
worker_stop = threading.Event()
worker_wake = threading.Event()
WEB_ROOT = Path(__file__).parent / "static"


def _purge_old_snapshots() -> None:
    """Keep the newest MAX DOM snapshots, drop the rest to bound disk usage."""
    snapshots = sorted(LOG_DIR.glob("max-*.html"), key=lambda path: path.stat().st_mtime)
    for stale in snapshots[:-MAX_SNAPSHOTS]:
        try:
            stale.unlink()
        except OSError:
            log.warning("Не удалось удалить старый снимок MAX %s", stale)


def _run_worker() -> None:
    while not worker_stop.is_set():
        worker_wake.wait(timeout=1.0)
        worker_wake.clear()
        if worker_stop.is_set():
            break
        try:
            with session_scope() as session:
                campaign = session.scalar(
                    select(Campaign).where(Campaign.status == "RUNNING").order_by(Campaign.id)
                )
                if campaign is None:
                    continue
                task = session.scalar(
                    select(DeliveryTask).where(
                        DeliveryTask.campaign_id == campaign.id,
                        DeliveryTask.status == "PENDING",
                    )
                )
                if task is None:
                    if session.scalar(
                        select(DeliveryTask).where(
                            DeliveryTask.campaign_id == campaign.id,
                            DeliveryTask.status == "REQUIRES_REVIEW",
                        )
                    ):
                        campaign.status = "WAITING_REVIEW"
                    else:
                        campaign.status = "COMPLETED"
                        campaign.finished_at = now()
                    session.commit()
                    continue
                task.status = "PROCESSING"
                task.started_at = now()
                task.attempts += 1
                if campaign.mode == "manual":
                    task.status = "REQUIRES_REVIEW"
                    task.error = "Ожидается подтверждение оператора"
                    task.requires_review = True
                task.finished_at = now()
                session.commit()
        except SQLAlchemyError:
            log.exception("delivery-worker iteration failed")
            time.sleep(0.5)


def start_worker() -> None:
    global worker_thread
    with worker_lock:
        if worker_thread is None or not worker_thread.is_alive():
            worker_stop.clear()
            worker_thread = threading.Thread(
                target=_run_worker, daemon=True, name="delivery-worker"
            )
            worker_thread.start()


def stop_worker() -> None:
    worker_stop.set()
    worker_wake.set()
    thread = worker_thread
    if thread is not None and thread.is_alive():
        thread.join(timeout=2)


def init_db() -> None:
    """Create missing tables and apply migrations. Idempotent and safe to repeat."""
    Base.metadata.create_all(engine)
    apply_migrations(engine)


@asynccontextmanager
async def lifespan(_app: FastAPI) -> AsyncIterator[None]:
    init_db()
    _purge_old_snapshots()
    start_worker()
    log.info("Application started; MAX browser preparation is available through CDP")
    try:
        yield
    finally:
        stop_worker()


app = FastAPI(title="MAX Рассылки", version="1.0.0", lifespan=lifespan)

# Гарантируем готовую схему и без запуска lifespan (например, в тестах TestClient).
init_db()


@app.middleware("http")
async def request_logging(request: Request, call_next: Any) -> Any:
    if request.url.path == HEALTH_PATH:
        return await call_next(request)
    request_id = uuid.uuid4().hex[:12]
    started = time.monotonic()
    log.info("request.start id=%s method=%s path=%s", request_id, request.method, request.url.path)
    try:
        response = await call_next(request)
    except Exception:
        log.exception("request.error id=%s method=%s path=%s", request_id, request.method, request.url.path)
        raise
    duration_ms = round((time.monotonic() - started) * 1000)
    response.headers["X-Request-ID"] = request_id
    log.info(
        "request.end id=%s status=%s duration_ms=%s",
        request_id,
        response.status_code,
        duration_ms,
    )
    return response


@app.get("/", include_in_schema=False)
def web_app() -> FileResponse:
    return FileResponse(WEB_ROOT / "index.html")


UI_ASSETS: dict[str, str] = {
    "styles.css": "text/css; charset=utf-8",
    "app.js": "text/javascript; charset=utf-8",
}


@app.get("/static/{asset_name}", include_in_schema=False)
def ui_asset(asset_name: str) -> FileResponse:
    """Отдаёт файлы интерфейса с явным MIME-типом.

    StaticFiles/mimetypes зависят от реестра Windows, где .js может оказаться
    text/plain, а строгая проверка MIME в Chrome тогда отклонит стили и скрипты.
    """
    media_type = UI_ASSETS.get(asset_name)
    path = WEB_ROOT / asset_name
    if media_type is None or not path.is_file():
        raise HTTPException(404, "Файл интерфейса не найден")
    return FileResponse(path, media_type=media_type)


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {
        "status": "ok",
        "worker": bool(worker_thread and worker_thread.is_alive()),
        "integration": "browser_cdp",
        "mode": "manual",
    }


@app.post("/api/auth/setup")
def setup_login(data: LoginIn) -> dict[str, str]:
    with session_scope() as session:
        if session.get(Setting, "password_hash"):
            raise HTTPException(409, "Пароль уже настроен")
        session.add(Setting(key="password_hash", value=hash_password(data.password)))
        session.commit()
    return {"status": "configured"}


@app.post("/api/auth/login")
def login(_: bool = Depends(require_auth)) -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/clients")
def list_clients(
    q: str = "",
    opted_out: bool | None = None,
    school_id: int | None = None,
    class_id: int | None = None,
    _: bool = Depends(require_auth),
) -> list[dict[str, Any]]:
    with session_scope() as session:
        query = select(Client).order_by(Client.id.desc())
        if q:
            query = query.where(Client.name.ilike(f"%{q}%"))
        if opted_out is not None:
            query = query.where(Client.opted_out == opted_out)
        if school_id is not None:
            query = query.where(Client.school_id == school_id)
        if class_id is not None:
            query = query.where(Client.class_id == class_id)
        return [client_dict(c) for c in session.scalars(query).all()]


@app.get("/api/schools")
def list_schools(_: bool = Depends(require_auth)) -> list[dict[str, Any]]:
    with session_scope() as session:
        schools = session.scalars(select(School).order_by(School.name)).all()
        result = []
        for school in schools:
            classes = session.scalars(
                select(SchoolClass)
                .where(SchoolClass.school_id == school.id)
                .order_by(SchoolClass.name)
            ).all()
            class_data = []
            for item in classes:
                contacts = session.scalars(
                    select(Client)
                    .where(Client.class_id == item.id)
                    .order_by(Client.role, Client.name)
                ).all()
                class_data.append(
                    {
                        "id": item.id,
                        "name": item.name,
                        "contacts": [
                            {
                                "id": contact.id,
                                "name": contact.name,
                                "phone": contact.phone,
                                "role": contact.role,
                                "opted_out": contact.opted_out,
                            }
                            for contact in contacts
                        ],
                    }
                )
            result.append({"id": school.id, "name": school.name, "classes": class_data})
        return result


@app.post("/api/classes")
def create_class(data: ClassIn, _: bool = Depends(require_auth)) -> dict[str, Any]:
    school_name = data.school_name.strip()
    class_name = data.class_name.strip()
    with session_scope() as session:
        school = session.scalar(select(School).where(School.name == school_name))
        if school is None:
            school = School(name=school_name)
            session.add(school)
            session.flush()
        existing = session.scalar(
            select(SchoolClass).where(
                SchoolClass.school_id == school.id, SchoolClass.name == class_name
            )
        )
        if existing:
            return {"school_id": school.id, "class_id": existing.id, "status": "existing"}
        item = SchoolClass(school_id=school.id, name=class_name)
        session.add(item)
        try:
            session.commit()
        except IntegrityError as exc:
            session.rollback()
            raise HTTPException(409, "Такой класс уже существует в этой школе") from exc
        return {"school_id": school.id, "class_id": item.id, "status": "created"}


@app.patch("/api/schools/{school_id}")
def update_school(
    school_id: int, data: SchoolUpdateIn, _: bool = Depends(require_auth)
) -> dict[str, Any]:
    with session_scope() as session:
        school = session.get(School, school_id)
        if school is None:
            raise HTTPException(404, "Школа не найдена")
        name = data.name.strip()
        duplicate = session.scalar(
            select(School).where(School.name == name, School.id != school_id)
        )
        if duplicate:
            raise HTTPException(409, "Школа с таким названием уже существует")
        school.name = name
        session.commit()
        return {"id": school.id, "name": school.name}


@app.delete("/api/schools/{school_id}")
def delete_school(school_id: int, _: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        school = session.get(School, school_id)
        if school is None:
            raise HTTPException(404, "Школа не найдена")
        class_ids = session.scalars(
            select(SchoolClass.id).where(SchoolClass.school_id == school_id)
        ).all()
        client_filter = Client.school_id == school_id
        if class_ids:
            client_filter = client_filter | Client.class_id.in_(class_ids)
        client_ids = session.scalars(select(Client.id).where(client_filter)).all()
        if client_ids:
            session.query(DeliveryTask).filter(
                DeliveryTask.client_id.in_(client_ids)
            ).delete(synchronize_session=False)
            session.query(Client).filter(Client.id.in_(client_ids)).delete(
                synchronize_session=False
            )
        if class_ids:
            session.query(SchoolClass).filter(SchoolClass.id.in_(class_ids)).delete(
                synchronize_session=False
            )
        session.delete(school)
        session.commit()
    return {"status": "deleted"}



@app.patch("/api/classes/{class_id}")
def update_class(
    class_id: int, data: ClassUpdateIn, _: bool = Depends(require_auth)
) -> dict[str, Any]:
    with session_scope() as session:
        item = session.get(SchoolClass, class_id)
        if item is None:
            raise HTTPException(404, "Класс не найден")
        duplicate = session.scalar(
            select(SchoolClass).where(
                SchoolClass.school_id == item.school_id,
                SchoolClass.name == data.name.strip(),
                SchoolClass.id != class_id,
            )
        )
        if duplicate:
            raise HTTPException(409, "Такой класс уже существует в этой школе")
        item.name = data.name.strip()
        session.commit()
        return {"id": item.id, "school_id": item.school_id, "name": item.name}


@app.delete("/api/classes/{class_id}")
def delete_class(class_id: int, _: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        item = session.get(SchoolClass, class_id)
        if item is None:
            raise HTTPException(404, "Класс не найден")
        client_ids = session.scalars(
            select(Client.id).where(Client.class_id == class_id)
        ).all()
        if client_ids:
            session.query(DeliveryTask).filter(
                DeliveryTask.client_id.in_(client_ids)
            ).delete(synchronize_session=False)
            session.query(Client).filter(Client.id.in_(client_ids)).delete(
                synchronize_session=False
            )
        session.delete(item)
        session.commit()
    return {"status": "deleted"}


def _validate_client_links(session: Any, school_id: int | None, class_id: int | None) -> None:
    if school_id is not None and session.get(School, school_id) is None:
        raise HTTPException(400, "Школа не найдена")
    if class_id is not None and session.get(SchoolClass, class_id) is None:
        raise HTTPException(400, "Класс не найден")
    if class_id is not None and school_id is not None:
        class_item = session.get(SchoolClass, class_id)
        if class_item and class_item.school_id != school_id:
            raise HTTPException(400, "Класс не относится к выбранной школе")


def _find_client_duplicate(session: Any, values: dict[str, Any]) -> Client | None:
    from sqlalchemy import or_

    conditions = [
        getattr(Client, key) == values[key]
        for key in DUPLICATE_FIELDS
        if values.get(key)
    ]
    if not conditions:
        return None
    return session.scalar(select(Client).where(or_(*conditions)))


@app.post("/api/clients")
def create_client(data: ClientIn, _: bool = Depends(require_auth)) -> dict[str, Any]:
    if data.role not in {"KR", "RK"}:
        raise HTTPException(400, "Роль должна быть KR или RK")
    with session_scope() as session:
        _validate_client_links(session, data.school_id, data.class_id)
        if _find_client_duplicate(session, data.model_dump()) is not None:
            raise HTTPException(409, "Клиент с таким идентификатором уже существует")
        payload = data.model_dump(exclude={"custom_fields"})
        payload["phone"] = normalize_phone(payload.get("phone"))
        c = Client(**payload, custom_fields=json.dumps(data.custom_fields, ensure_ascii=False))
        session.add(c)
        session.commit()
        session.refresh(c)
        return client_dict(c)


@app.patch("/api/clients/{client_id}")
def update_client(client_id: int, data: ClientIn, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        c = session.get(Client, client_id)
        if c is None:
            raise HTTPException(404, "Клиент не найден")
        _validate_client_links(session, data.school_id, data.class_id)
        for key, value in data.model_dump(exclude={"custom_fields"}).items():
            setattr(c, key, normalize_phone(value) if key == "phone" else value)
        c.custom_fields = json.dumps(data.custom_fields, ensure_ascii=False)
        session.commit()
        return client_dict(c)


@app.delete("/api/clients/{client_id}")
def delete_client(client_id: int, _: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        c = session.get(Client, client_id)
        if c is None:
            raise HTTPException(404, "Клиент не найден")
        session.query(DeliveryTask).filter(DeliveryTask.client_id == client_id).delete(
            synchronize_session=False
        )
        session.delete(c)
        session.commit()
    return {"status": "deleted"}


@app.post("/api/clients/delete")
def delete_clients(data: ClientDeleteIn, _: bool = Depends(require_auth)) -> dict[str, int]:
    client_ids = list(dict.fromkeys(data.client_ids))
    with session_scope() as session:
        existing_ids = session.scalars(
            select(Client.id).where(Client.id.in_(client_ids))
        ).all()
        if not existing_ids:
            raise HTTPException(404, "Выбранные контакты не найдены")
        session.query(DeliveryTask).filter(
            DeliveryTask.client_id.in_(existing_ids)
        ).delete(synchronize_session=False)
        deleted = session.query(Client).filter(
            Client.id.in_(existing_ids)
        ).delete(synchronize_session=False)
        session.commit()
    return {"deleted": deleted}



def _campaign_template(campaign: Campaign, role: str) -> str:
    if role == "RK":
        template = campaign.message_rk or campaign.message_kr or campaign.message
    elif role == "KR":
        template = campaign.message_kr or campaign.message or campaign.message_rk
    else:
        template = campaign.message or campaign.message_kr or campaign.message_rk
    return template or ""


def campaign_dict(campaign: Campaign, tasks: list[DeliveryTask]) -> dict[str, Any]:
    statuses: dict[str, int] = {}
    for task in tasks:
        statuses[task.status] = statuses.get(task.status, 0) + 1
    return {
        "id": campaign.id,
        "name": campaign.name,
        "description": campaign.description,
        "message": campaign.message,
        "message_kr": campaign.message_kr,
        "message_rk": campaign.message_rk,
        "school_id": campaign.school_id,
        "mode": campaign.mode,
        "status": campaign.status,
        "created_at": campaign.created_at,
        "started_at": campaign.started_at,
        "finished_at": campaign.finished_at,
        "stop_reason": campaign.stop_reason,
        "task_count": len(tasks),
        "status_counts": statuses,
    }


def task_dict(task: DeliveryTask, contact: Client, school_name: str, class_name: str) -> dict[str, Any]:
    return {
        "id": task.id,
        "campaign_id": task.campaign_id,
        "client_id": task.client_id,
        "client": contact.name,
        "phone": contact.phone,
        "chat_link": contact.chat_link,
        "max_user_id": contact.max_user_id,
        "username": contact.username,
        "role": contact.role,
        "recipient_role": task.recipient_role,
        "school": school_name,
        "class_name": class_name,
        "message": task.rendered_message,
        "status": task.status,
        "attempts": task.attempts,
        "error": task.error,
        "requires_review": task.requires_review,
        "message_id": task.message_id,
        "opened_at": task.opened_at,
        "operator_confirmed_at": task.operator_confirmed_at,
        "operator_result": task.operator_result,
        "browser_prepared_at": task.browser_prepared_at,
        "browser_error": task.browser_error,
        "created_at": task.created_at,
        "started_at": task.started_at,
        "finished_at": task.finished_at,
    }


def _load_campaign_tasks(session: Any, campaign_id: int) -> list[dict[str, Any]]:
    rows = session.execute(
        select(DeliveryTask, Client, School.name, SchoolClass.name)
        .join(Client, Client.id == DeliveryTask.client_id)
        .outerjoin(SchoolClass, SchoolClass.id == Client.class_id)
        .outerjoin(School, School.id == Client.school_id)
        .where(DeliveryTask.campaign_id == campaign_id)
        .order_by(DeliveryTask.id)
    ).all()
    return [
        task_dict(task, contact, school_name or "", class_name or "")
        for task, contact, school_name, class_name in rows
    ]


@app.get("/api/campaigns")
def list_campaigns(_: bool = Depends(require_auth)) -> list[dict[str, Any]]:
    with session_scope() as session:
        campaigns = session.scalars(select(Campaign).order_by(Campaign.id.desc())).all()
        result = []
        for campaign in campaigns:
            tasks = session.scalars(
                select(DeliveryTask).where(DeliveryTask.campaign_id == campaign.id)
            ).all()
            result.append(campaign_dict(campaign, list(tasks)))
        return result


@app.post("/api/campaigns")
def create_campaign(data: CampaignIn, _: bool = Depends(require_auth)) -> dict[str, Any]:
    client_ids = list(dict.fromkeys(data.client_ids))
    if not client_ids:
        raise HTTPException(400, "Не выбраны получатели")
    with session_scope() as session:
        contacts = session.scalars(select(Client).where(Client.id.in_(client_ids))).all()
        if len(contacts) != len(client_ids):
            raise HTTPException(404, "Некоторые получатели не найдены")
        if data.school_id is not None and session.get(School, data.school_id) is None:
            raise HTTPException(400, "Школа не найдена")
        campaign = Campaign(
            name=data.name.strip(),
            description=data.description,
            message=data.message,
            message_kr=data.message_kr,
            message_rk=data.message_rk,
            school_id=data.school_id,
            mode="manual",
            status="DRAFT",
        )
        session.add(campaign)
        session.flush()
        for contact in contacts:
            template = _campaign_template(campaign, contact.role)
            session.add(
                DeliveryTask(
                    campaign_id=campaign.id,
                    client_id=contact.id,
                    recipient_role=contact.role,
                    rendered_message=render_message(template, contact),
                    status="PENDING",
                )
            )
        session.commit()
        session.refresh(campaign)
        return campaign_dict(campaign, [])



@app.get("/api/campaigns/{campaign_id}")
def get_campaign(campaign_id: int, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        campaign = session.get(Campaign, campaign_id)
        if campaign is None:
            raise HTTPException(404, "Кампания не найдена")
        tasks = session.scalars(
            select(DeliveryTask).where(DeliveryTask.campaign_id == campaign_id)
        ).all()
        return campaign_dict(campaign, list(tasks))


@app.get("/api/campaigns/{campaign_id}/tasks")
def campaign_tasks(campaign_id: int, _: bool = Depends(require_auth)) -> list[dict[str, Any]]:
    with session_scope() as session:
        return _load_campaign_tasks(session, campaign_id)


@app.post("/api/campaigns/{campaign_id}/start")
def start_campaign(campaign_id: int, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        campaign = session.get(Campaign, campaign_id)
        if campaign is None:
            raise HTTPException(404, "Кампания не найдена")
        if campaign.status == "RUNNING":
            return {"status": "running"}
        campaign.status = "RUNNING"
        campaign.started_at = now()
        campaign.finished_at = None
        session.commit()
    worker_wake.set()
    return {"status": "running"}


@app.post("/api/campaigns/{campaign_id}/stop")
def stop_campaign(
    campaign_id: int, reason: str = "", _: bool = Depends(require_auth)
) -> dict[str, Any]:
    with session_scope() as session:
        campaign = session.get(Campaign, campaign_id)
        if campaign is None:
            raise HTTPException(404, "Кампания не найдена")
        campaign.status = "STOPPED"
        campaign.finished_at = now()
        campaign.stop_reason = reason or "Остановлено оператором"
        session.commit()
    return {"status": "stopped"}


@app.delete("/api/campaigns/{campaign_id}")
def delete_campaign(campaign_id: int, _: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        campaign = session.get(Campaign, campaign_id)
        if campaign is None:
            raise HTTPException(404, "Кампания не найдена")
        session.query(DeliveryTask).filter(
            DeliveryTask.campaign_id == campaign_id
        ).delete(synchronize_session=False)
        session.delete(campaign)
        session.commit()
    return {"status": "deleted"}


@app.post("/api/campaigns/delete")
def delete_campaigns(data: CampaignDeleteIn, _: bool = Depends(require_auth)) -> dict[str, int]:
    campaign_ids = list(dict.fromkeys(data.campaign_ids))
    with session_scope() as session:
        existing_ids = session.scalars(
            select(Campaign.id).where(Campaign.id.in_(campaign_ids))
        ).all()
        if not existing_ids:
            raise HTTPException(404, "Выбранные кампании не найдены")
        session.query(DeliveryTask).filter(
            DeliveryTask.campaign_id.in_(existing_ids)
        ).delete(synchronize_session=False)
        deleted = session.query(Campaign).filter(
            Campaign.id.in_(existing_ids)
        ).delete(synchronize_session=False)
        session.commit()
    return {"deleted": deleted}



def _rows_from_upload(filename: str, raw: bytes) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        for merged_range in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            if min_col != max_col:
                continue
            value = sheet.cell(min_row, min_col).value
            if value is None:
                continue
            column_index = min_col - 1
            for row_index in range(min_row, max_row + 1):
                rows[row_index - 1][column_index] = value
        workbook.close()
        return rows
    text = raw.decode("utf-8-sig", errors="replace")
    counts = {candidate: text.count(candidate) for candidate in ",;\t|"}
    delimiter = max(counts, key=lambda candidate: counts[candidate]) if any(counts.values()) else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _build_preview(rows: list[list[Any]]) -> dict[str, Any]:
    def cell_text(cell: Any) -> str:
        if cell is None:
            return ""
        # Excel returns whole numbers as float (89099695363.0); show them as integers.
        if isinstance(cell, float) and cell.is_integer():
            return str(int(cell))
        return str(cell).strip()

    header_index = detect_header_row(rows)
    raw_headers = rows[header_index] if rows else []
    headers = [cell_text(cell) for cell in raw_headers]
    data_rows = rows[header_index + 1 :]
    preview_rows = []
    for row in data_rows:
        if all(cell_text(cell) == "" for cell in row):
            continue
        record = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            record[header] = "" if index >= len(row) else cell_text(row[index])
        preview_rows.append(record)
    return {
        "header_row": header_index + 1,
        "headers": headers,
        "rows": preview_rows,
        "total_rows": len(preview_rows),
    }


@app.post("/api/imports/preview")
async def import_preview(
    file: UploadFile = File(...), _: bool = Depends(require_auth)
) -> dict[str, Any]:
    raw = await file.read()
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Файл слишком большой (максимум 10 МБ)")
    try:
        rows = _rows_from_upload(file.filename or "upload.csv", raw)
    except Exception as exc:
        raise HTTPException(400, f"Не удалось прочитать файл: {exc}") from exc
    return _build_preview(rows)


@app.post("/api/imports/confirm")
def import_confirm(data: ImportConfirm, _: bool = Depends(require_auth)) -> dict[str, int]:
    created = updated = skipped = 0
    with session_scope() as session:
        for row in data.rows:
            values: dict[str, Any] = {}
            for header, field in data.mapping.items():
                if field in {"name", "phone", "max_user_id", "username", "chat_link"}:
                    values[field] = row.get(header)
            name = (values.get("name") or "").strip()
            if not name:
                skipped += 1
                continue
            values["phone"] = normalize_phone(values.get("phone"))
            duplicate = _find_client_duplicate(session, values)
            if duplicate is not None:
                if data.duplicate_mode == "update":
                    duplicate.name = name
                    duplicate.phone = values.get("phone") or duplicate.phone
                    duplicate.chat_link = values.get("chat_link") or duplicate.chat_link
                    duplicate.max_user_id = values.get("max_user_id") or duplicate.max_user_id
                    duplicate.username = values.get("username") or duplicate.username
                    updated += 1
                else:
                    skipped += 1
                continue
            session.add(
                Client(
                    name=name,
                    phone=values.get("phone"),
                    max_user_id=values.get("max_user_id"),
                    username=values.get("username"),
                    chat_link=values.get("chat_link"),
                    source="import",
                )
            )
            created += 1
        session.commit()
    return {"created": created, "updated": updated, "skipped": skipped}



def _resolve_school(session: Any, school_name: str) -> School:
    """Accept either a numeric school id or a human-readable school name."""
    value = (school_name or "").strip()
    if not value:
        raise HTTPException(400, "Не указано название школы")
    if value.isdigit():
        school = session.get(School, int(value))
        if school is not None:
            return school
    school = session.scalar(select(School).where(School.name == value))
    if school is None:
        school = School(name=value)
        session.add(school)
        session.flush()
    return school


def _resolve_class(session: Any, school_id: int, class_name: str) -> SchoolClass:
    value = (class_name or "").strip()
    if not value:
        value = "Без класса"
    item = session.scalar(
        select(SchoolClass).where(
            SchoolClass.school_id == school_id, SchoolClass.name == value
        )
    )
    if item is None:
        item = SchoolClass(school_id=school_id, name=value)
        session.add(item)
        session.flush()
    return item


def _normalize_import_header(value: Any) -> str:
    return "".join(str(value or "").casefold().replace("ё", "е").split())


def _import_row_value(row: dict[str, Any], header: str | None) -> Any:
    if header is None:
        return None
    if header in row:
        return row[header]
    normalized_header = _normalize_import_header(header)
    return next(
        (value for key, value in row.items() if _normalize_import_header(key) == normalized_header),
        None,
    )


@app.post("/api/imports/school-contacts")
def import_school_contacts(
    data: ContactImport, _: bool = Depends(require_auth)
) -> dict[str, int]:
    if not data.school_name:
        raise HTTPException(400, "Не указано название школы")
    created = updated = skipped = 0
    headers_by_field: dict[str, str] = {}
    for header, field in data.mapping.items():
        if not field:
            continue
        # При нескольких столбцах с одним полем берём первый по порядку в файле.
        headers_by_field.setdefault(field, header)
    class_header = headers_by_field.get("class_name")
    last_class_name = ""
    with session_scope() as session:
        school = _resolve_school(session, data.school_name)
        for row in data.rows:
            class_name = str(_import_row_value(row, class_header) or "").strip() if class_header else ""
            if class_name:
                last_class_name = class_name
            elif last_class_name:
                # Excel often stores the class only in the first row of a merged cell.
                class_name = last_class_name
            item = _resolve_class(session, school.id, class_name)
            for role, name_field, phone_field in (
                ("KR", "kr_name", "kr_phone"),
                ("RK", "rk_name", "rk_phone"),
            ):
                name_header = headers_by_field.get(name_field)
                if name_header is None:
                    continue
                name = str(_import_row_value(row, name_header) or "").strip()
                if not name:
                    continue
                phone_header = headers_by_field.get(phone_field)
                phone = normalize_phone(
                    str(_import_row_value(row, phone_header) or "").strip() if phone_header else None
                )
                duplicate = session.scalar(
                    select(Client).where(
                        Client.school_id == school.id,
                        Client.class_id == item.id,
                        Client.name == name,
                        Client.role == role,
                    )
                )
                if duplicate is None:
                    duplicate = session.scalar(
                        select(Client)
                        .outerjoin(SchoolClass, Client.class_id == SchoolClass.id)
                        .where(
                            Client.school_id == school.id,
                            Client.name == name,
                            Client.role == role,
                            (Client.class_id.is_(None)) | (SchoolClass.name == "Без класса"),
                        )
                    )
                if duplicate is not None:
                    if data.duplicate_mode == "update" and duplicate.class_id != item.id:
                        duplicate.class_id = item.id
                        duplicate.phone = phone or duplicate.phone
                        updated += 1
                    else:
                        skipped += 1
                    continue
                session.add(
                    Client(
                        name=name,
                        phone=phone,
                        role=role,
                        school_id=school.id,
                        class_id=item.id,
                        source="import",
                    )
                )
                created += 1
        session.commit()
    return {"created": created, "updated": updated, "skipped": skipped}


@app.patch("/api/settings")
def update_settings(data: SettingsIn, _: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        setting = session.get(Setting, "delivery_mode")
        if setting is None:
            setting = Setting(key="delivery_mode", value="manual")
            session.add(setting)
        setting.value = "manual"
        session.commit()
    return {"mode": "manual"}


@app.get("/api/settings")
def get_settings(_: bool = Depends(require_auth)) -> dict[str, str]:
    with session_scope() as session:
        setting = session.get(Setting, "delivery_mode")
        return {"mode": setting.value if setting else "manual"}



def _task_and_contact(session: Any, task_id: int) -> tuple[DeliveryTask, Client]:
    task = session.get(DeliveryTask, task_id)
    if task is None:
        raise HTTPException(404, "Задача не найдена")
    contact = session.get(Client, task.client_id)
    if contact is None:
        raise HTTPException(404, "Клиент задачи не найден")
    return task, contact


def _contact_identity(contact: Client) -> str:
    return contact.max_user_id or contact.chat_link or contact.username or ""


@app.post("/api/tasks/{task_id}/prepare-max")
def prepare_max_task(task_id: int, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        task, contact = _task_and_contact(session, task_id)
        if not contact.phone:
            raise HTTPException(502, "У контакта не указан номер телефона для поиска в MAX")
        client_name = contact.name
        identity = contact.phone
        message = task.rendered_message
    try:
        result = prepare_message(client_name, identity, message)
    except MaxBrowserError as exc:
        with session_scope() as session:
            failed = session.get(DeliveryTask, task_id)
            if failed is not None:
                if "не найден" in str(exc).casefold():
                    failed.status = "CONTACT_NOT_FOUND"
                    failed.browser_error = None
                else:
                    failed.browser_error = str(exc)
                session.commit()
        raise HTTPException(502, str(exc)) from exc
    with session_scope() as session:
        task = session.get(DeliveryTask, task_id)
        if task is not None:
            task.browser_prepared_at = now()
            task.browser_error = None
            session.commit()
    return {
        "status": result.status,
        "page_url": result.page_url,
        "client": result.client_name,
    }


@app.post("/api/tasks/{task_id}/open-max")
def open_max_task(task_id: int, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        task, contact = _task_and_contact(session, task_id)
        client_name = contact.name
        identity = contact.phone or _contact_identity(contact)
    try:
        result = open_chat(client_name, identity)
    except MaxBrowserError as exc:
        raise HTTPException(502, str(exc)) from exc
    with session_scope() as session:
        task = session.get(DeliveryTask, task_id)
        if task is not None:
            task.opened_at = now()
            session.commit()
    return {"status": result.status, "page_url": result.page_url, "client": result.client_name}


@app.post("/api/tasks/{task_id}/opened")
def mark_task_opened(task_id: int, _: bool = Depends(require_auth)) -> dict[str, Any]:
    with session_scope() as session:
        task, _contact = _task_and_contact(session, task_id)
        task.opened_at = now()
        session.commit()
        return {"id": task.id, "opened_at": task.opened_at}


@app.patch("/api/tasks/{task_id}/status")
def update_task_status(
    task_id: int, data: TaskStatusIn, _: bool = Depends(require_auth)
) -> dict[str, Any]:
    with session_scope() as session:
        task, _contact = _task_and_contact(session, task_id)
        task.status = data.status
        if data.status in {"SENT", "FAILED", "CANCELLED"}:
            task.finished_at = now()
        session.commit()
        return {"id": task.id, "status": task.status}



@app.post("/api/tasks/{task_id}/confirm")
def confirm_manual_task(
    task_id: int, data: ManualResultIn, _: bool = Depends(require_auth)
) -> dict[str, Any]:
    with session_scope() as session:
        task, _contact = _task_and_contact(session, task_id)
        mapping = {"sent": "SENT", "failed": "FAILED", "unknown": "UNKNOWN"}
        task.operator_result = data.result
        task.operator_confirmed_at = now()
        task.status = mapping[data.result]
        task.requires_review = False
        if data.note:
            task.error = data.note
        if task.status in {"SENT", "FAILED"}:
            task.finished_at = now()
        session.commit()
        return {"id": task.id, "status": task.status, "operator_result": task.operator_result}


@app.get("/api/reports/{campaign_id}/export")
def export_report(campaign_id: int, _: bool = Depends(require_auth)) -> StreamingResponse:
    with session_scope() as session:
        campaign = session.get(Campaign, campaign_id)
        if campaign is None:
            raise HTTPException(404, "Кампания не найдена")
        header = [
            "Кампания",
            "ФИО",
            "Телефон",
            "Роль",
            "Школа",
            "Класс",
            "Статус",
            "Результат оператора",
            "Попыток",
            "Ошибка",
            "Подготовлено",
            "Подтверждено",
        ]
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
        writer.writerow(header)
        for item in _load_campaign_tasks(session, campaign_id):
            writer.writerow(
                [
                    csv_safe(campaign.name),
                    csv_safe(item["client"]),
                    csv_safe(item["phone"]),
                    csv_safe(item["recipient_role"]),
                    csv_safe(item["school"]),
                    csv_safe(item["class_name"]),
                    csv_safe(item["status"]),
                    csv_safe(item["operator_result"]),
                    csv_safe(item["attempts"]),
                    csv_safe(item["error"]),
                    csv_safe(item["browser_prepared_at"]),
                    csv_safe(item["operator_confirmed_at"]),
                ]
            )
        payload = "\ufeff" + buffer.getvalue()
    filename = f"campaign-{campaign_id}-report.csv"
    return StreamingResponse(
        iter([payload.encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

