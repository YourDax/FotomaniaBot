import io

from fastapi.testclient import TestClient
from uuid import uuid4

from openpyxl import Workbook
from app.main import app
from app import main
from app.max_browser import (
    MaxBrowserError,
    PreparationResult,
    _is_phone_not_found_text,
    _max_phone_digits,
)

client = TestClient(app)


def test_max_phone_digits_removes_russian_prefixes():
    assert _max_phone_digits("+79176370371") == "9176370371"
    assert _max_phone_digits("+7 999 123-45-67") == "9991234567"
    assert _max_phone_digits("8 (999) 123-45-67") == "9991234567"
    assert _max_phone_digits("9991234567") == "9991234567"


def test_max_phone_not_found_messages_are_detected():
    assert _is_phone_not_found_text(
        "Не нашли номер +79266231948\nОтправьте ссылку‑приглашение на чат с вами"
    )
    assert _is_phone_not_found_text("Контакт не найден")
    assert not _is_phone_not_found_text("Чат открыт")


def test_health():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["integration"] == "browser_cdp"


def test_school_and_class_crud():
    school = client.post(
        "/api/classes",
        json={"school_name": f"CRUD школа {uuid4()}", "class_name": "5А"},
    ).json()
    school_id, class_id = school["school_id"], school["class_id"]
    assert client.patch(f"/api/schools/{school_id}", json={"name": "Переименованная школа"}).status_code == 200
    assert client.patch(f"/api/classes/{class_id}", json={"name": "5Б"}).status_code == 200
    assert client.delete(f"/api/classes/{class_id}").status_code == 200
    assert client.delete(f"/api/schools/{school_id}").status_code == 200


def test_deleting_school_and_class_cascades_only_their_contacts():
    first = client.post(
        "/api/classes",
        json={"school_name": f"Школа каскад 1 {uuid4()}", "class_name": "1А"},
    ).json()
    second = client.post(
        "/api/classes",
        json={"school_name": f"Школа каскад 2 {uuid4()}", "class_name": "2А"},
    ).json()
    first_contact = client.post(
        "/api/clients",
        json={
            "name": "Контакт первой школы",
            "phone": "79990000001",
            "school_id": first["school_id"],
            "class_id": first["class_id"],
        },
    ).json()
    second_contact = client.post(
        "/api/clients",
        json={
            "name": "Контакт второй школы",
            "phone": "79990000002",
            "school_id": second["school_id"],
            "class_id": second["class_id"],
        },
    ).json()
    assert client.delete(f"/api/classes/{first['class_id']}").status_code == 200
    remaining = {item["id"] for item in client.get("/api/clients").json()}
    assert first_contact["id"] not in remaining
    assert second_contact["id"] in remaining
    assert client.delete(f"/api/schools/{second['school_id']}").status_code == 200
    assert second_contact["id"] not in {item["id"] for item in client.get("/api/clients").json()}


def test_client_and_campaign_flow():
    name = "Тестовый клиент"
    response = client.post("/api/clients", json={"name": name, "chat_link": f"https://max.ru/{uuid4()}"})
    assert response.status_code == 200
    client_id = response.json()["id"]
    response = client.post("/api/campaigns", json={"name": "Тест", "message": "Здравствуйте, {name}!", "client_ids": [client_id]})
    assert response.status_code == 200
    campaign_id = response.json()["id"]
    assert client.get("/api/campaigns").json()[0]["mode"] == "manual"
    assert client.post(f"/api/campaigns/{campaign_id}/start").status_code == 200


def test_campaign_tasks_include_class_without_message_text():
    school_name = f"Класс рассылки {uuid4()}"
    school = client.post(
        "/api/classes",
        json={"school_name": school_name, "class_name": "7Б"},
    ).json()
    contact = client.post(
        "/api/clients",
        json={
            "name": "Контакт рассылки",
            "phone": "79990000009",
            "school_id": school["school_id"],
            "class_id": school["class_id"],
        },
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": "Рассылка класса", "message": "Скрытый текст", "client_ids": [contact["id"]]},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert task["class_name"] == "7Б"
    assert task["school"] == school_name
    assert task["message"] == "Скрытый текст"


def test_campaign_can_be_deleted_with_tasks():
    contact = client.post("/api/clients", json={"name": f"Удаляемый {uuid4()}"}).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": "Удаляемая рассылка", "message": "Текст", "client_ids": [contact["id"]]},
    ).json()
    response = client.delete(f"/api/campaigns/{campaign['id']}")
    assert response.status_code == 200
    assert all(item["id"] != campaign["id"] for item in client.get("/api/campaigns").json())
    assert client.get(f"/api/campaigns/{campaign['id']}/tasks").status_code == 200
    assert client.get(f"/api/campaigns/{campaign['id']}/tasks").json() == []


def test_campaigns_can_be_deleted_in_bulk():
    campaign_ids = []
    for index in range(2):
        contact = client.post("/api/clients", json={"name": f"Bulk {uuid4()}"}).json()
        campaign = client.post(
            "/api/campaigns",
            json={"name": f"Bulk campaign {index}", "message": "Текст", "client_ids": [contact["id"]]},
        ).json()
        campaign_ids.append(campaign["id"])
    response = client.post("/api/campaigns/delete", json={"campaign_ids": campaign_ids})
    assert response.status_code == 200
    assert response.json()["deleted"] == 2


def test_clients_can_be_deleted_in_bulk_with_delivery_tasks():
    first = client.post("/api/clients", json={"name": f"Массовое удаление 1 {uuid4()}"}).json()
    second = client.post("/api/clients", json={"name": f"Массовое удаление 2 {uuid4()}"}).json()
    campaign = client.post(
        "/api/campaigns",
        json={
            "name": f"Рассылка для массового удаления {uuid4()}",
            "message": "Текст",
            "client_ids": [first["id"]],
        },
    ).json()
    response = client.post(
        "/api/clients/delete",
        json={"client_ids": [first["id"], first["id"], second["id"]]},
    )
    assert response.status_code == 200
    assert response.json()["deleted"] == 2
    assert client.get(f"/api/campaigns/{campaign['id']}/tasks").json() == []
    remaining = {item["id"] for item in client.get("/api/clients").json()}
    assert first["id"] not in remaining
    assert second["id"] not in remaining


def test_single_client_delete_cascades_delivery_tasks():
    contact = client.post("/api/clients", json={"name": f"Одиночное удаление {uuid4()}"}).json()
    campaign = client.post(
        "/api/campaigns",
        json={
            "name": f"Рассылка для одиночного удаления {uuid4()}",
            "message": "Текст",
            "client_ids": [contact["id"]],
        },
    ).json()
    assert client.get(f"/api/campaigns/{campaign['id']}/tasks").json()
    assert client.delete(f"/api/clients/{contact['id']}").status_code == 200
    assert client.get(f"/api/campaigns/{campaign['id']}/tasks").json() == []
    assert contact["id"] not in {item["id"] for item in client.get("/api/clients").json()}


def test_school_contacts_import_uses_school_name_not_id():
    school_name = f"Школа импорта имени {uuid4()}"
    created = client.post(
        "/api/classes",
        json={"school_name": school_name, "class_name": "1А"},
    ).json()
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": [
                {
                    "Класс": "5А",
                    "ФИО КР": "КР Импорт",
                    "Телефон КР": "79991110001",
                    "ФИО РК": "РК Импорт",
                    "Телефон РК": "79991110002",
                }
            ],
            "mapping": {
                "Класс": "class_name",
                "ФИО КР": "kr_name",
                "Телефон КР": "kr_phone",
                "ФИО РК": "rk_name",
                "Телефон РК": "rk_phone",
            },
            "school_name": school_name,
        },
    )
    assert response.status_code == 200
    schools = {item["id"]: item["name"] for item in client.get("/api/schools").json()}
    assert schools[created["school_id"]] == school_name
    assert str(created["school_id"]) not in schools.values()
    contacts = [
        item
        for item in client.get("/api/clients").json()
        if item["school_id"] == created["school_id"] and item["name"] in {"КР Импорт", "РК Импорт"}
    ]
    assert len(contacts) == 2


def test_school_contacts_import_resolves_numeric_school_id():
    school_name = f"Школа по id {uuid4()}"
    created = client.post(
        "/api/classes",
        json={"school_name": school_name, "class_name": "2А"},
    ).json()
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": [
                {
                    "Класс": "6Б",
                    "ФИО КР": "КР По ID",
                    "Телефон КР": "79992220001",
                    "ФИО РК": "РК По ID",
                    "Телефон РК": "79992220002",
                }
            ],
            "mapping": {
                "Класс": "class_name",
                "ФИО КР": "kr_name",
                "Телефон КР": "kr_phone",
                "ФИО РК": "rk_name",
                "Телефон РК": "rk_phone",
            },
            "school_name": str(created["school_id"]),
        },
    )
    assert response.status_code == 200
    schools = client.get("/api/schools").json()
    assert not any(item["name"] == str(created["school_id"]) for item in schools)
    contacts = [
        item
        for item in client.get("/api/clients").json()
        if item["school_id"] == created["school_id"] and item["name"] in {"КР По ID", "РК По ID"}
    ]
    assert len(contacts) == 2


def test_auth_blocks_without_credentials_when_password_configured():
    password = f"secret-{uuid4().hex}"
    assert client.post("/api/auth/setup", json={"password": password}).status_code == 200
    try:
        assert client.get("/api/clients").status_code == 401
        ok = client.get("/api/clients", auth=("operator", password))
        assert ok.status_code == 200
    finally:
        with main.db() as session:
            setting = session.get(main.Setting, "password_hash")
            if setting is not None:
                session.delete(setting)
                session.commit()


def test_csv_import_and_settings():
    suffix = uuid4().hex
    csv_data = f"Имя,Телефон,Ссылка\nИмпортированный клиент,{suffix},https://max.ru/{suffix}\n"
    response = client.post(
        "/api/imports/preview",
        files={"file": ("clients.csv", csv_data.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200
    preview = response.json()
    assert preview["total_rows"] == 1
    response = client.post(
        "/api/imports/confirm",
        json={
            "rows": preview["rows"],
            "mapping": {"Имя": "name", "Телефон": "phone", "Ссылка": "chat_link"},
            "duplicate_mode": "skip",
        },
    )
    assert response.status_code == 200
    assert response.json()["created"] == 1
    assert client.patch("/api/settings", json={"mode": "manual"}).status_code == 200


def test_school_contacts_import_inherits_merged_class_and_uses_mapped_header():
    suffix = uuid4().hex[:8]
    school_name = f"Школа объединённых ячеек {suffix}"
    rows = [
        {"Название класса": "5А", "ФИО КР": f"КР 5А {suffix}"},
        {"Название класса": "", "ФИО КР": f"КР 5А второй {suffix}"},
        {"Название класса": "6Б", "ФИО КР": f"КР 6Б {suffix}"},
        {"Название класса": "", "ФИО КР": f"КР 6Б второй {suffix}"},
    ]
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": rows,
            "mapping": {"Название класса": "class_name", "ФИО КР": "kr_name"},
            "school_name": school_name,
        },
    )

    assert response.status_code == 200
    assert response.json() == {"created": 4, "updated": 0, "skipped": 0}
    imported = {
        item["name"]: item
        for item in client.get("/api/clients").json()
        if item["name"].endswith(suffix)
    }
    assert len(imported) == 4
    class_ids = {name: contact["class_id"] for name, contact in imported.items()}
    assert class_ids[f"КР 5А {suffix}"] == class_ids[f"КР 5А второй {suffix}"]
    assert class_ids[f"КР 6Б {suffix}"] == class_ids[f"КР 6Б второй {suffix}"]
    assert class_ids[f"КР 5А {suffix}"] != class_ids[f"КР 6Б {suffix}"]


def test_school_contacts_import_update_reassigns_existing_no_class_contacts():
    suffix = uuid4().hex[:8]
    school_name = f"Школа восстановления классов {suffix}"
    placeholder = client.post(
        "/api/classes",
        json={"school_name": school_name, "class_name": "Без класса"},
    ).json()
    old_contact = client.post(
        "/api/clients",
        json={
            "name": f"КР без класса {suffix}",
            "phone": "79990000000",
            "school_id": placeholder["school_id"],
            "class_id": placeholder["class_id"],
            "role": "KR",
        },
    ).json()

    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": [{"Класс": "2В", "ФИО КР": f"КР без класса {suffix}", "Телефон КР": "79990000001"}],
            "mapping": {"Класс": "class_name", "ФИО КР": "kr_name", "Телефон КР": "kr_phone"},
            "school_name": school_name,
            "duplicate_mode": "update",
        },
    )

    assert response.status_code == 200
    assert response.json() == {"created": 0, "updated": 1, "skipped": 0}
    saved_contact = next(
        item for item in client.get("/api/clients").json() if item["id"] == old_contact["id"]
    )
    assert saved_contact["school_id"] == placeholder["school_id"]
    assert saved_contact["class_id"] != placeholder["class_id"]
    classes = next(
        item["classes"] for item in client.get("/api/schools").json()
        if item["id"] == placeholder["school_id"]
    )
    assigned_class = next(item for item in classes if item["id"] == saved_contact["class_id"])
    assert assigned_class["name"] == "2В"
    assert assigned_class["contacts"][0]["id"] == old_contact["id"]



def test_school_contacts_import_normalizes_mapped_headers_with_newlines():
    suffix = uuid4().hex[:8]
    school_name = f"Школа заголовков с переносами {suffix}"
    rows = [
        {"Класс": "5А", "ФИО КР": f"КР 5А {suffix}"},
        {"Класс": "", "ФИО КР": f"КР 5А второй {suffix}"},
        {"Класс": "6Б", "ФИО КР": f"КР 6Б {suffix}"},
    ]
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": rows,
            "mapping": {"Класс\n": "class_name", "ФИО\nКР": "kr_name"},
            "school_name": school_name,
        },
    )

    assert response.status_code == 200
    assert response.json() == {"created": 3, "updated": 0, "skipped": 0}
    imported = {
        item["name"]: item
        for item in client.get("/api/clients").json()
        if item["name"].endswith(suffix)
    }
    assert imported[f"КР 5А {suffix}"]["class_id"] == imported[f"КР 5А второй {suffix}"]["class_id"]
    assert imported[f"КР 5А {suffix}"]["class_id"] != imported[f"КР 6Б {suffix}"]["class_id"]


def test_school_contacts_import_creates_separate_roles():
    rows = [{"Класс": "5А", "ФИО КР": "Иванов Иван", "Телефон КР": "79990000001", "ФИО РК": "Петрова Анна", "Телефон РК": "79990000002"}]
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": rows,
            "mapping": {"Класс": "class_name", "ФИО КР": "kr_name", "Телефон КР": "kr_phone", "ФИО РК": "rk_name", "Телефон РК": "rk_phone"},
            "school_name": "Школа №1",
        },
    )
    assert response.status_code == 200
    contacts = client.get("/api/clients").json()
    school_contacts = [item for item in contacts if item["class_id"] and item["role"] in {"KR", "RK"}]
    assert {item["role"] for item in school_contacts} == {"KR", "RK"}


def test_excel_preview_imports_contacts_into_distinct_classes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Школа: Тестовая"])
    sheet.append([None])
    sheet.append(["No", "Класс", "Кабинет", "ФИО КР", "Телефон КР", "ФИО РК", "Телефон РК"])
    sheet.append([1, "5А", None, None, "79990000001", None, None])
    sheet.append([2, None, None, None, None, None, None])
    sheet.append([3, "6Б", None, None, "79990000002", None, None])
    sheet.cell(row=4, column=4, value="КР 5А")
    sheet.cell(row=5, column=4, value="КР 5А второй")
    sheet.cell(row=6, column=4, value="КР 6Б")
    sheet.merge_cells("B4:B5")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    suffix = uuid4().hex[:8]

    preview_response = client.post(
        "/api/imports/preview",
        files={"file": ("school.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert preview["header_row"] == 3
    assert [row["Класс"] for row in preview["rows"]] == ["5А", "5А", "6Б"]
    import_response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": preview["rows"],
            "mapping": {"Класс": "class_name", "ФИО КР": "kr_name", "Телефон КР": "kr_phone"},
            "school_name": f"Школа XLSX {suffix}",
        },
    )

    assert import_response.status_code == 200
    assert import_response.json() == {"created": 3, "updated": 0, "skipped": 0}
    school = next(
        item for item in client.get("/api/schools").json() if item["name"] == f"Школа XLSX {suffix}"
    )
    classes = {item["name"]: item for item in school["classes"]}
    assert set(classes) == {"5А", "6Б"}
    assert {contact["name"] for contact in classes["5А"]["contacts"]} == {"КР 5А", "КР 5А второй"}
    assert {contact["name"] for contact in classes["6Б"]["contacts"]} == {"КР 6Б"}


def test_import_detects_header_row_not_at_top():
    csv_data = "\n\nСписок класса\nКласс,ФИО КР,Телефон КР,ФИО РК,Телефон РК\n5Б,Иванов Иван,79990000003,Петрова Анна,79990000004\n"
    response = client.post(
        "/api/imports/preview",
        files={"file": ("school.csv", csv_data.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["header_row"] == 4
    assert payload["headers"] == ["Класс", "ФИО КР", "Телефон КР", "ФИО РК", "Телефон РК"]
    assert payload["rows"][0]["Класс"] == "5Б"


def test_school_contacts_import_uses_first_class_column_when_others_are_empty():
    suffix = uuid4().hex[:8]
    school_name = f"Школа дубля класса {suffix}"
    response = client.post(
        "/api/imports/school-contacts",
        json={
            "rows": [
                {
                    "Класс": "5А",
                    "Доп.Информация класса": "",
                    "Снято на Первоклассник Виньетка (1 кл)": "",
                    "ФИО КР": f"КР дубля {suffix}",
                    "Телефон КР": "79990000005",
                }
            ],
            "mapping": {
                "Класс": "class_name",
                "Доп.Информация класса": "class_name",
                "Снято на Первоклассник Виньетка (1 кл)": "class_name",
                "ФИО КР": "kr_name",
                "Телефон КР": "kr_phone",
            },
            "school_name": school_name,
        },
    )

    assert response.status_code == 200
    assert response.json() == {"created": 1, "updated": 0, "skipped": 0}
    school = next(item for item in client.get("/api/schools").json() if item["name"] == school_name)
    assert {item["name"] for item in school["classes"]} == {"5А"}


def test_manual_task_open_and_confirmation_are_persisted():
    suffix = uuid4().hex
    school = client.post("/api/classes", json={"school_name": f"Школа ручного режима {suffix}", "class_name": f"6А-{suffix[:4]}"}).json()
    contact = client.post(
        "/api/clients",
        json={"name": "КР для проверки", "phone": suffix, "chat_link": f"https://max.ru/manual-check-{suffix}", "school_id": school["school_id"], "class_id": school["class_id"], "role": "KR"},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": "Ручная проверка", "school_id": school["school_id"], "message_kr": "Текст для {name}", "client_ids": [contact["id"]], "mode": "manual"},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert client.post(f"/api/tasks/{task['id']}/opened").status_code == 200
    assert client.post(f"/api/tasks/{task['id']}/confirm", json={"result": "sent"}).status_code == 200
    saved = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert saved["operator_result"] == "sent"
    assert saved["opened_at"] is not None


def test_task_status_can_be_updated_manually():
    suffix = uuid4().hex
    contact = client.post(
        "/api/clients",
        json={"name": f"Статус {suffix}", "phone": f"7999{suffix[:7]}"},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": f"Статус рассылки {suffix}", "message": "Текст", "client_ids": [contact["id"]]},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    response = client.patch(f"/api/tasks/{task['id']}/status", json={"status": "CONTACT_NOT_FOUND"})
    assert response.status_code == 200
    saved = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert saved["status"] == "CONTACT_NOT_FOUND"


def test_campaign_can_use_selected_contacts_from_multiple_schools():
    first = client.post(
        "/api/classes",
        json={"school_name": f"Первая школа {uuid4()}", "class_name": "1А"},
    ).json()
    second = client.post(
        "/api/classes",
        json={"school_name": f"Вторая школа {uuid4()}", "class_name": "2А"},
    ).json()
    first_contact = client.post(
        "/api/clients",
        json={"name": "Первый получатель", "phone": "79990000101", "school_id": first["school_id"], "class_id": first["class_id"]},
    ).json()
    second_contact = client.post(
        "/api/clients",
        json={"name": "Второй получатель", "phone": "79990000102", "school_id": second["school_id"], "class_id": second["class_id"]},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={
            "name": "Несколько школ",
            "message": "Общее сообщение",
            "client_ids": [first_contact["id"], second_contact["id"]],
        },
    )
    assert campaign.status_code == 200
    tasks = client.get(f"/api/campaigns/{campaign.json()['id']}/tasks").json()
    assert {task["client"] for task in tasks} == {"Первый получатель", "Второй получатель"}


def test_prepare_max_task_persists_verified_browser_draft(monkeypatch):
    suffix = uuid4().hex
    contact = client.post(
        "/api/clients",
        json={"name": f"MAX тест {suffix}", "phone": f"+7999{suffix[:7]}"},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={
            "name": f"MAX подготовка {suffix}",
            "message": "Здравствуйте, {name}!",
            "client_ids": [contact["id"]],
            "mode": "manual",
        },
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    client.post(f"/api/campaigns/{campaign['id']}/start")

    monkeypatch.setattr(
        main,
        "prepare_message",
        lambda name, identity, message: PreparationResult(
            "DRAFT_PREPARED", "https://max.ru/chats/test", name, message
        ),
    )
    response = client.post(f"/api/tasks/{task['id']}/prepare-max")
    assert response.status_code == 200
    assert response.json()["status"] == "DRAFT_PREPARED"
    saved = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert saved["browser_prepared_at"] is not None
    assert saved["browser_error"] is None


def test_prepare_max_requires_phone_number():
    contact = client.post("/api/clients", json={"name": f"Без телефона {uuid4()}"}).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": "Телефон обязателен", "message": "Текст", "client_ids": [contact["id"]]},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    response = client.post(f"/api/tasks/{task['id']}/prepare-max")
    assert response.status_code == 502
    assert "номер телефона" in response.json()["detail"]


def test_open_max_task_opens_chat_without_inserting_text(monkeypatch):
    suffix = uuid4().hex
    contact = client.post(
        "/api/clients",
        json={"name": f"Открыть чат {suffix}", "phone": f"+7999{suffix[:7]}"},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": f"Открытие чата {suffix}", "message": "Не вставлять", "client_ids": [contact["id"]]},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    monkeypatch.setattr(
        main,
        "open_chat",
        lambda name, identity: PreparationResult(
            "CHAT_OPENED", "https://web.max.ru/chat", name, ""
        ),
    )
    response = client.post(f"/api/tasks/{task['id']}/open-max")
    assert response.status_code == 200
    assert response.json()["status"] == "CHAT_OPENED"
    saved = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert saved["opened_at"] is not None
    assert saved["browser_prepared_at"] is None


def test_prepare_max_contact_not_found_is_a_status(monkeypatch):
    suffix = uuid4().hex
    contact = client.post(
        "/api/clients",
        json={"name": f"Не найден {suffix}", "phone": f"+7999{suffix[:7]}"},
    ).json()
    campaign = client.post(
        "/api/campaigns",
        json={"name": f"Поиск {suffix}", "message": "Текст", "client_ids": [contact["id"]]},
    ).json()
    task = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    client.post(f"/api/campaigns/{campaign['id']}/start")

    def fail_prepare(name, identity, message):
        raise MaxBrowserError("Контакт не найден в MAX по указанному номеру телефона.")

    monkeypatch.setattr(main, "prepare_message", fail_prepare)
    response = client.post(f"/api/tasks/{task['id']}/prepare-max")
    assert response.status_code == 502
    saved = client.get(f"/api/campaigns/{campaign['id']}/tasks").json()[0]
    assert saved["status"] == "CONTACT_NOT_FOUND"
    assert saved["browser_error"] is None
